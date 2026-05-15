"""Normalize spreadsheets into sheet-aware Markdown and row-level source blocks.

The XLSX loader reads cached cell values, creates one section per worksheet,
keeps header context, and emits row blocks with common policy metadata fields
for filtering during retrieval.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.ingest.loaders.models import NormalizedSource, SourceBlock


def normalize_xlsx(source_path: Path) -> NormalizedSource:
    # Read formulas as cached values because RAG needs visible business data.
    workbook = load_workbook(filename=source_path, read_only=True, data_only=True)
    markdown_sections: list[str] = []
    blocks: list[SourceBlock] = []
    for sheet in workbook.worksheets:
        section = f"Sheet: {sheet.title}"
        rows = [_clean_row(row) for row in sheet.iter_rows(values_only=True)]
        non_empty_rows = [row for row in rows if any(row)]

        markdown_sections.append(f"## {section}")
        blocks.append(
            SourceBlock(
                text=f"## {section}",
                block_type="heading",
                section_path=section,
                sheet=sheet.title,
            )
        )

        if not non_empty_rows:
            markdown_sections.append("")
            continue

        headers = _build_headers(non_empty_rows[0])
        header_text = _headers_text(headers)
        markdown_sections.append(_markdown_row(headers))
        blocks.append(
            SourceBlock(
                text=header_text,
                block_type="table_header",
                section_path=section,
                sheet=sheet.title,
                metadata={
                    "row_number": 1,
                    "column_headers": headers,
                },
            )
        )

        for row_index, row in enumerate(non_empty_rows[1:], start=2):
            if not any(row):
                continue

            markdown_sections.append(_markdown_row(row))
            row_values = _row_values(headers, row)
            blocks.append(
                SourceBlock(
                    text=_row_block_text(row_values),
                    block_type="table_row",
                    section_path=section,
                    sheet=sheet.title,
                    metadata={
                        "row_number": row_index,
                        "column_headers": headers,
                        "row_values": row_values,
                        **_filter_metadata_from_row(row_values),
                    },
                )
            )

        # If the sheet only has one populated row, still keep it as inspectable data.
        if len(non_empty_rows) == 1:
            row_values = _row_values(headers, non_empty_rows[0])
            blocks.append(
                SourceBlock(
                    text=_row_block_text(row_values),
                    block_type="table_row",
                    section_path=section,
                    sheet=sheet.title,
                    metadata={
                        "row_number": 1,
                        "column_headers": headers,
                        "row_values": row_values,
                        **_filter_metadata_from_row(row_values),
                    },
                )
            )
        markdown_sections.append("")

    return NormalizedSource(
        markdown_text="\n".join(markdown_sections).strip(),
        blocks=blocks,
        extraction_method="openpyxl_row_level",
    )


def _clean_row(row: tuple[Any, ...]) -> list[str]:
    return ["" if value is None else str(value).strip() for value in row]


def _build_headers(header_row: list[str]) -> list[str]:
    headers: list[str] = []
    for index, value in enumerate(header_row, start=1):
        headers.append(value or f"column_{index}")
    return headers


def _row_values(headers: list[str], row: list[str]) -> dict[str, str]:
    values: dict[str, str] = {}
    for index, header in enumerate(headers):
        values[header] = row[index] if index < len(row) else ""
    return values


def _headers_text(headers: list[str]) -> str:
    return "Columns: " + ", ".join(headers)


def _row_block_text(row_values: dict[str, str]) -> str:
    lines = [
        f"{header}: {value}" if value else f"{header}:"
        for header, value in row_values.items()
    ]
    return "\n".join(lines)


def _markdown_row(values: list[str]) -> str:
    return "| " + " | ".join(values) + " |"


def _filter_metadata_from_row(row_values: dict[str, str]) -> dict[str, str]:
    # These common policy dimensions are useful for Chroma metadata filters.
    metadata: dict[str, str] = {}
    for source_key, metadata_key in {
        "country_code": "country_code",
        "country": "country",
        "city": "city",
        "category": "expense_category",
        "expense_category": "expense_category",
        "currency": "currency",
    }.items():
        value = row_values.get(source_key)
        if value:
            metadata[metadata_key] = value
    return metadata
