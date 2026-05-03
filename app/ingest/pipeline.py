from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from hashlib import sha1
from pathlib import Path
from typing import Iterable

from app.core.config import Settings
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {".pdf", ".html", ".htm", ".xlsx", ".png", ".jpg", ".jpeg"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}


@dataclass(frozen=True)
class IngestedDocument:
    doc_id: str
    source_path: str
    output_markdown_path: str
    doc_type: str
    title: str
    markdown_text: str
    extraction_method: str
    extraction_warning: str | None


def discover_source_files(raw_data_dir: Path) -> list[Path]:
    return sorted(
        path
        for path in raw_data_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )


def ingest_sources(settings: Settings) -> list[IngestedDocument]:
    settings.markdown_dir.mkdir(parents=True, exist_ok=True)
    settings.processed_data_dir.mkdir(parents=True, exist_ok=True)

    ingested_documents: list[IngestedDocument] = []
    for source_path in discover_source_files(settings.raw_data_dir):
        document = normalize_source(source_path, settings)
        write_markdown(document)
        ingested_documents.append(document)

    write_manifest(settings.processed_data_dir / "ingest_manifest.jsonl", ingested_documents)
    return ingested_documents


def normalize_source(source_path: Path, settings: Settings) -> IngestedDocument:
    suffix = source_path.suffix.lower()
    markdown_text, extraction_method, extraction_warning = _normalize_content(
        source_path
    )
    doc_id = _build_doc_id(source_path)
    doc_type = _infer_doc_type(suffix)
    title = source_path.stem.replace("_", " ").replace("-", " ").title()
    output_path = settings.markdown_dir / f"{doc_id}.md"

    header = [
        f"# {title}",
        "",
        f"- doc_id: `{doc_id}`",
        f"- source_path: `{source_path.as_posix()}`",
        f"- doc_type: `{doc_type}`",
        f"- extraction_method: `{extraction_method}`",
    ]
    if extraction_warning:
        header.append(f"- extraction_warning: `{extraction_warning}`")

    markdown = "\n".join(header + ["", "## Content", "", markdown_text.strip(), ""]).strip()

    return IngestedDocument(
        doc_id=doc_id,
        source_path=source_path.as_posix(),
        output_markdown_path=output_path.as_posix(),
        doc_type=doc_type,
        title=title,
        markdown_text=markdown,
        extraction_method=extraction_method,
        extraction_warning=extraction_warning,
    )


def write_markdown(document: IngestedDocument) -> None:
    output_path = Path(document.output_markdown_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(document.markdown_text + "\n", encoding="utf-8")


def write_manifest(output_path: Path, documents: Iterable[IngestedDocument]) -> None:
    rows = [json.dumps(asdict(document), ensure_ascii=True) for document in documents]
    output_path.write_text("\n".join(rows) + ("\n" if rows else ""), encoding="utf-8")


def _normalize_content(source_path: Path) -> tuple[str, str, str | None]:
    suffix = source_path.suffix.lower()
    if suffix in {".html", ".htm"}:
        return _normalize_html(source_path)
    if suffix == ".xlsx":
        return _normalize_xlsx(source_path)
    if suffix == ".pdf":
        return _normalize_pdf(source_path)
    if suffix in IMAGE_SUFFIXES:
        return _normalize_image(source_path)
    return ("Unsupported file type.", "unsupported", "unsupported_suffix")


def _normalize_html(source_path: Path) -> tuple[str, str, str | None]:
    html_text = source_path.read_text(encoding="utf-8")
    soup = BeautifulSoup(html_text, "html.parser")
    for tag_name in ("script", "style"):
        for tag in soup.find_all(tag_name):
            tag.decompose()

    # Future: evaluate Unstructured or html2text if this loses useful structure.
    return (_html_soup_to_markdown(soup), "beautifulsoup4", None)


def _normalize_xlsx(source_path: Path) -> tuple[str, str, str | None]:
    workbook = load_workbook(filename=source_path, read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets:
        sections.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                sections.append("| " + " | ".join(values) + " |")
        sections.append("")

    # Future: use UnstructuredExcelLoader if table structure gets more complex.
    return ("\n".join(sections).strip(), "openpyxl", None)


def _normalize_pdf(source_path: Path) -> tuple[str, str, str | None]:
    reader = PdfReader(str(source_path))
    sections: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = _clean_text(page.extract_text() or "")
        sections.append(f"## Page {index}")
        sections.append(page_text or "[No text extracted from this page]")
        sections.append("")

    # Future: escalate to Unstructured/PyMuPDF/OCR when quality checks flag weak output.
    return ("\n".join(sections).strip(), "pypdf", None)


def _normalize_image(source_path: Path) -> tuple[str, str, str | None]:
    # Future: add OCR or a vision model pass for image-heavy policy sources.
    return (
        "\n".join(
            [
                f"Image source: `{source_path.name}`",
                "",
                "Image extraction is not implemented yet.",
            ]
        ),
        "image_placeholder",
        "image_extraction_not_implemented",
    )


def _build_doc_id(source_path: Path) -> str:
    digest = sha1(source_path.as_posix().encode("utf-8")).hexdigest()[:10]
    slug = re.sub(r"[^a-z0-9]+", "-", source_path.stem.lower()).strip("-")
    return f"{slug}-{digest}"


def _infer_doc_type(suffix: str) -> str:
    if suffix == ".pdf":
        return "pdf"
    if suffix in {".html", ".htm"}:
        return "html"
    if suffix == ".xlsx":
        return "xlsx"
    if suffix in IMAGE_SUFFIXES:
        return "image"
    return "unknown"


def _clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _html_soup_to_markdown(soup: BeautifulSoup) -> str:
    body = soup.body or soup
    blocks = _render_block_nodes(body)
    cleaned_blocks = [block.strip() for block in blocks if block and block.strip()]
    return "\n\n".join(cleaned_blocks)


def _render_block_nodes(node) -> list[str]:
    blocks: list[str] = []
    for child in getattr(node, "children", []):
        name = getattr(child, "name", None)
        if name is None:
            text = _clean_text(str(child))
            if text:
                blocks.append(text)
            continue

        if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = int(name[1])
            text = _collect_inline_text(child)
            if text:
                blocks.append(f"{'#' * level} {text}")
            continue

        if name == "p":
            text = _collect_inline_text(child)
            if text:
                blocks.append(text)
            continue

        if name in {"ul", "ol"}:
            blocks.extend(_render_list(child, ordered=name == "ol"))
            continue

        if name == "table":
            table_markdown = _render_table(child)
            if table_markdown:
                blocks.append(table_markdown)
            continue

        if name in {"section", "article", "main", "header", "footer", "nav", "div"}:
            blocks.extend(_render_block_nodes(child))
            continue

        text = _collect_inline_text(child)
        if text:
            blocks.append(text)

    return blocks


def _render_list(list_node, ordered: bool) -> list[str]:
    items: list[str] = []
    index = 1
    for item in list_node.find_all("li", recursive=False):
        text = _collect_inline_text(item)
        if not text:
            continue
        prefix = f"{index}. " if ordered else "- "
        items.append(f"{prefix}{text}")
        index += 1
    return items


def _render_table(table_node) -> str:
    rows: list[list[str]] = []
    for row in table_node.find_all("tr"):
        cells = row.find_all(["th", "td"])
        values = [_collect_inline_text(cell) for cell in cells]
        if any(values):
            rows.append(values)

    if not rows:
        return ""

    width = max(len(row) for row in rows)
    padded_rows = [row + [""] * (width - len(row)) for row in rows]
    header = padded_rows[0]
    separator = ["---"] * width
    body = padded_rows[1:]

    markdown_rows = [
        f"| {' | '.join(header)} |",
        f"| {' | '.join(separator)} |",
    ]
    markdown_rows.extend(f"| {' | '.join(row)} |" for row in body)
    return "\n".join(markdown_rows)


def _collect_inline_text(node) -> str:
    text_parts: list[str] = []
    for part in node.stripped_strings:
        text_parts.append(_clean_text(str(part)))
    return " ".join(part for part in text_parts if part)
