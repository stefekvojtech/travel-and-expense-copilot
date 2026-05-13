"""Convert currencies for claim evaluation with local rates before API fallback.

This module provides a deterministic currency conversion helper for future claim
tools. It should be used as a secondary option: policy retrieval and company
Finance rates remain the preferred source of reimbursement exchange rates. When
the local workbook cannot provide the requested conversion, the tool can fall
back to Frankfurter's free exchange-rate API.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Callable, Literal
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import urlopen

from openpyxl import load_workbook

from app.core.config import Settings


DEFAULT_FRANKFURTER_BASE_URL = "https://api.frankfurter.dev"
DEFAULT_TIMEOUT_SECONDS = 10

TOOL_USAGE_PROMPT = """Use convert_currency only as a secondary tool.

Preferred order for reimbursement exchange rates:
1. Use retrieved policy evidence or the company Finance exchange-rate workbook.
2. If the company rate is unavailable, use convert_currency with the external
   fallback enabled.
3. Do not use user-provided exchange rates for reimbursement decisions because
   they can be incorrect or manipulated.

The tool performs conversion only. It does not decide whether a claim is
reimbursable, whether a cap applies, or whether approval is required.
"""

JsonGetter = Callable[[str, int], dict[str, Any]]


class CurrencyConversionError(RuntimeError):
    """Raised when currency conversion cannot be completed."""


@dataclass(frozen=True)
class CurrencyConversion:
    """Structured output from the currency conversion tool."""

    amount: Decimal
    from_currency: str
    to_currency: str
    converted_amount: Decimal
    exchange_rate: Decimal
    source: Literal["identity", "company_exchange_rates", "frankfurter"]
    source_detail: str
    rate_date: str | None
    warnings: list[str]


@dataclass(frozen=True)
class WorkbookRate:
    """One local Finance exchange-rate row."""

    month: str
    currency: str
    eur_rate: Decimal
    source: str | None


def convert_currency(
    settings: Settings,
    amount: Decimal | int | float | str,
    from_currency: str,
    to_currency: str = "EUR",
    *,
    rate_date: date | str | None = None,
    allow_external_fallback: bool = True,
    frankfurter_base_url: str = DEFAULT_FRANKFURTER_BASE_URL,
    timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS,
    http_get_json: JsonGetter | None = None,
) -> CurrencyConversion:
    """Convert money using company rates first and Frankfurter only as fallback.

    The local company source is `data/raw/per_diem_caps.xlsx`, sheet
    `ExchangeRates`. Those rates are preferred because reimbursement policy says
    Finance monthly rates should be used. Frankfurter is used only when the
    workbook has no usable rate for the requested currency pair.
    """
    normalized_amount = _to_decimal(amount)
    source_currency = _normalize_currency(from_currency)
    target_currency = _normalize_currency(to_currency)
    requested_date = _normalize_rate_date(rate_date)

    if source_currency == target_currency:
        return CurrencyConversion(
            amount=normalized_amount,
            from_currency=source_currency,
            to_currency=target_currency,
            converted_amount=_round_money(normalized_amount),
            exchange_rate=Decimal("1"),
            source="identity",
            source_detail="No conversion needed because currencies match.",
            rate_date=requested_date,
            warnings=[],
        )

    workbook_rates = _load_company_exchange_rates(
        settings.raw_data_dir / "per_diem_caps.xlsx"
    )
    company_conversion = _convert_with_company_rates(
        amount=normalized_amount,
        from_currency=source_currency,
        to_currency=target_currency,
        requested_date=requested_date,
        workbook_rates=workbook_rates,
    )
    if company_conversion is not None:
        return company_conversion

    if not allow_external_fallback:
        raise CurrencyConversionError(
            "Company Finance exchange rates did not contain a usable conversion "
            f"from {source_currency} to {target_currency}, and external fallback "
            "is disabled."
        )

    return _convert_with_frankfurter(
        amount=normalized_amount,
        from_currency=source_currency,
        to_currency=target_currency,
        rate_date=requested_date,
        base_url=frankfurter_base_url,
        timeout_seconds=timeout_seconds,
        http_get_json=http_get_json or _http_get_json,
    )


def _load_company_exchange_rates(workbook_path: Path) -> list[WorkbookRate]:
    if not workbook_path.exists():
        return []

    workbook = load_workbook(filename=workbook_path, read_only=True, data_only=True)
    try:
        if "ExchangeRates" not in workbook.sheetnames:
            return []

        sheet = workbook["ExchangeRates"]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    rates: list[WorkbookRate] = []
    for row in rows[1:]:
        values = {
            headers[index]: value
            for index, value in enumerate(row)
            if index < len(headers) and headers[index]
        }
        month = values.get("month")
        currency = values.get("currency")
        eur_rate = values.get("eur_rate")
        if month in (None, "") or currency in (None, "") or eur_rate in (None, ""):
            continue
        rates.append(
            WorkbookRate(
                month=str(month).strip(),
                currency=_normalize_currency(str(currency)),
                eur_rate=_to_decimal(eur_rate),
                source=str(values.get("source")).strip()
                if values.get("source") not in (None, "")
                else None,
            )
        )
    return rates


def _convert_with_company_rates(
    *,
    amount: Decimal,
    from_currency: str,
    to_currency: str,
    requested_date: str | None,
    workbook_rates: list[WorkbookRate],
) -> CurrencyConversion | None:
    rates_by_currency = _select_company_rates(
        workbook_rates,
        requested_month=requested_date[:7] if requested_date else None,
    )
    if not rates_by_currency:
        return None

    source_to_eur = _eur_rate_for_currency(from_currency, rates_by_currency)
    target_to_eur = _eur_rate_for_currency(to_currency, rates_by_currency)
    if source_to_eur is None or target_to_eur is None:
        return None

    exchange_rate = source_to_eur / target_to_eur
    converted_amount = _round_money(amount * exchange_rate)
    source_months = sorted({rate.month for rate in rates_by_currency.values()})
    return CurrencyConversion(
        amount=amount,
        from_currency=from_currency,
        to_currency=to_currency,
        converted_amount=converted_amount,
        exchange_rate=exchange_rate,
        source="company_exchange_rates",
        source_detail="data/raw/per_diem_caps.xlsx: ExchangeRates sheet",
        rate_date=source_months[-1] if source_months else None,
        warnings=[],
    )


def _select_company_rates(
    rates: list[WorkbookRate],
    *,
    requested_month: str | None,
) -> dict[str, WorkbookRate]:
    if requested_month is not None:
        return {
            rate.currency: rate
            for rate in rates
            if rate.month == requested_month
        }

    latest_by_currency: dict[str, WorkbookRate] = {}
    for rate in sorted(rates, key=lambda item: item.month):
        latest_by_currency[rate.currency] = rate
    return latest_by_currency


def _eur_rate_for_currency(
    currency: str,
    rates_by_currency: dict[str, WorkbookRate],
) -> Decimal | None:
    if currency == "EUR":
        return Decimal("1")
    rate = rates_by_currency.get(currency)
    return rate.eur_rate if rate is not None else None


def _convert_with_frankfurter(
    *,
    amount: Decimal,
    from_currency: str,
    to_currency: str,
    rate_date: str | None,
    base_url: str,
    timeout_seconds: int,
    http_get_json: JsonGetter,
) -> CurrencyConversion:
    path = f"{base_url.rstrip('/')}/v2/rate/{from_currency}/{to_currency}"
    query = urlencode({"date": rate_date}) if rate_date is not None else ""
    url = f"{path}?{query}" if query else path
    payload = http_get_json(url, timeout_seconds)
    raw_rate = payload.get("rate")
    if raw_rate is None:
        raise CurrencyConversionError(
            "Frankfurter response did not include a `rate` field."
        )

    exchange_rate = _to_decimal(raw_rate)
    return CurrencyConversion(
        amount=amount,
        from_currency=from_currency,
        to_currency=to_currency,
        converted_amount=_round_money(amount * exchange_rate),
        exchange_rate=exchange_rate,
        source="frankfurter",
        source_detail="Frankfurter free exchange-rate API",
        rate_date=str(payload.get("date")) if payload.get("date") else rate_date,
        warnings=[
            "Used external Frankfurter fallback because company Finance rates "
            "did not contain the requested conversion."
        ],
    )


def _http_get_json(url: str, timeout_seconds: int) -> dict[str, Any]:
    try:
        with urlopen(url, timeout=timeout_seconds) as response:
            data = response.read().decode("utf-8")
    except HTTPError as exc:
        raise CurrencyConversionError(
            f"Frankfurter returned HTTP {exc.code} for currency conversion."
        ) from exc
    except URLError as exc:
        raise CurrencyConversionError(
            f"Could not reach Frankfurter for currency conversion: {exc.reason}"
        ) from exc

    decoded = json.loads(data)
    if not isinstance(decoded, dict):
        raise CurrencyConversionError(
            "Frankfurter returned an unexpected non-object JSON response."
        )
    return decoded


def _normalize_currency(currency: str) -> str:
    normalized = currency.strip().upper()
    if len(normalized) != 3 or not normalized.isalpha():
        raise CurrencyConversionError(
            f"Currency must be a three-letter ISO code, got {currency!r}."
        )
    return normalized


def _normalize_rate_date(value: date | str | None) -> str | None:
    if value is None:
        return None
    if isinstance(value, date):
        return value.isoformat()

    text = value.strip()
    if len(text) == 7:
        return f"{text}-01"
    if len(text) == 10:
        return text
    raise CurrencyConversionError(
        "rate_date must be either YYYY-MM-DD, YYYY-MM, or a datetime.date."
    )


def _to_decimal(value: Decimal | int | float | str) -> Decimal:
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _round_money(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
